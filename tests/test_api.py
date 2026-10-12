import io
from datetime import datetime, time, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, inspect, select

from backend.config import settings
from backend.database import SessionLocal, engine
from backend.models import Entity, FieldVisit, FleetIntegration, User, VehiclePosition, VisitType
from services.telematics.sync import sync_integration
from tests.conftest import auth_headers


def _yesterday() -> str:
    return (datetime.now(settings.tz).date() - timedelta(days=1)).isoformat()


async def test_database_is_migrated_to_head():
    async with engine.connect() as conn:
        tables = await conn.run_sync(lambda sync: set(inspect(sync).get_table_names()))
        version = (await conn.exec_driver_sql("SELECT version_num FROM alembic_version")).scalar_one()
    assert {"company", "company_locations", "fleet_integrations", "vehicle_positions", "daily_route_reports"} <= tables
    assert version == "0002"


@pytest.mark.parametrize(
    ("username", "path", "expected"),
    [
        ("pazarlamaci", "/api/fleet/live", 403),
        ("surucu", "/api/reports/archive", 403),
        ("lojistik", "/api/fleet/live", 200),
        ("lojistik", "/api/reports/archive", 200),
        ("lojistik", "/api/fleet/providers", 403),
        ("pazarlamaci", "/api/company/locations", 200),
        ("admin", "/api/fleet/providers", 200),
    ],
)
async def test_role_permissions(client, username, path, expected):
    headers = await auth_headers(client, username)
    response = await client.get(path, headers=headers)
    assert response.status_code == expected, response.text


async def test_dashboard_counts_visits_by_local_day(client):
    headers = await auth_headers(client, "admin")
    before = (await client.get("/api/dashboard/summary", headers=headers)).json()["completed_visits_today"]

    # 00:30 in Istanbul is still the previous day in UTC.
    after_midnight = datetime.combine(datetime.now(settings.tz).date(), time(0, 30), settings.tz)
    async with SessionLocal() as db:
        user = await db.scalar(select(User).where(User.username == "pazarlamaci"))
        entity = await db.scalar(select(Entity).limit(1))
        db.add(
            FieldVisit(
                user_id=user.id,
                entity_id=entity.id,
                visit_type=VisitType.marketing,
                check_in_time=after_midnight - timedelta(minutes=20),
                check_out_time=after_midnight,
            )
        )
        await db.commit()

    after = (await client.get("/api/dashboard/summary", headers=headers)).json()["completed_visits_today"]
    assert after == before + 1


async def test_non_admin_cannot_change_company_or_integrations(client):
    headers = await auth_headers(client, "lojistik")
    company = await client.put("/api/company", json={"name": "Başka Firma"}, headers=headers)
    created = await client.post(
        "/api/fleet/integrations",
        json={"provider": "simulator", "name": "Yetkisiz", "credentials": {}},
        headers=headers,
    )
    assert company.status_code == 403
    assert created.status_code == 403


async def test_integration_secrets_are_never_returned(client):
    headers = await auth_headers(client, "admin")
    secret = "arvento-gizli-parola-123"
    created = await client.post(
        "/api/fleet/integrations",
        json={
            "provider": "arvento",
            "name": "Arvento deneme",
            "credentials": {"endpoint": "https://arvento.example", "username": "demo", "password": secret},
        },
        headers=headers,
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert secret not in created.text
    assert body["configured_fields"] == ["endpoint", "password", "username"]

    listed = await client.get("/api/fleet/integrations", headers=headers)
    assert secret not in listed.text

    async with SessionLocal() as db:
        stored = await db.get(FleetIntegration, body["id"])
        assert secret not in stored.credentials_encrypted

    tested = await client.post(f"/api/fleet/integrations/{body['id']}/test", headers=headers)
    assert tested.status_code == 200
    assert tested.json()["ok"] is False
    assert "henüz tamamlanmadı" in tested.json()["message"]

    updated = await client.put(
        f"/api/fleet/integrations/{body['id']}",
        json={"name": "Arvento ana hesap", "credentials": {"customer_code": "C-1"}},
        headers=headers,
    )
    assert updated.status_code == 200
    assert updated.json()["configured_fields"] == ["customer_code", "endpoint", "password", "username"]

    missing = await client.post(
        "/api/fleet/integrations",
        json={"provider": "mobiliz", "name": "Eksik", "credentials": {"endpoint": "https://m.example"}},
        headers=headers,
    )
    assert missing.status_code == 400

    deleted = await client.delete(f"/api/fleet/integrations/{body['id']}", headers=headers)
    assert deleted.status_code == 200


async def test_sync_deduplicates_positions(client):
    headers = await auth_headers(client, "admin")
    integrations = (await client.get("/api/fleet/integrations", headers=headers)).json()
    simulator = next(item for item in integrations if item["provider"] == "simulator")

    response = await client.post(f"/api/fleet/integrations/{simulator['id']}/sync", headers=headers)
    assert response.status_code == 200

    # The simulator only drives during shift hours, so pin the clock to yesterday noon.
    noon = datetime.combine(datetime.now(settings.tz).date() - timedelta(days=1), time(12), settings.tz)
    async with SessionLocal() as db:
        integration = await db.get(FleetIntegration, simulator["id"])
        integration.last_sync_at = None
        first = await sync_integration(db, integration, noon)
        await sync_integration(db, integration, noon + timedelta(minutes=5))
    assert first.unmatched_ids == ["SIM-02", "SIM-03"]

    async with SessionLocal() as db:
        total = await db.scalar(select(func.count()).select_from(VehiclePosition))
        distinct = await db.scalar(
            select(func.count()).select_from(
                select(VehiclePosition.vehicle_id, VehiclePosition.recorded_at).distinct().subquery()
            )
        )
    assert total == distinct


async def test_track_report_and_archive_flow(client):
    headers = await auth_headers(client, "lojistik")
    live = (await client.get("/api/fleet/live", headers=headers)).json()
    van = next(vehicle for vehicle in live if vehicle["external_id"] == "SIM-01")
    day = _yesterday()

    track = await client.get(f"/api/fleet/vehicles/{van['id']}/track?date={day}", headers=headers)
    assert track.status_code == 200
    body = track.json()
    assert len(body["points"]) > 100
    assert body["stops"] and body["summary"]["distance_km"] > 10
    assert "Tavas Kooperatif Market" in body["missed_planned"]
    assert any(stop["planned"] for stop in body["stops"])

    pdf = await client.get(f"/api/reports/route-day?vehicle_id={van['id']}&date={day}&format=pdf", headers=headers)
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    assert "filename*=UTF-8''" in pdf.headers["content-disposition"]

    xlsx = await client.get(f"/api/reports/route-day?vehicle_id={van['id']}&date={day}&format=xlsx", headers=headers)
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Özet", "Duraklar", "Konumlar"]
    assert workbook["Konumlar"].max_row > 100

    fleet = await client.get(f"/api/reports/fleet-day?date={day}", headers=headers)
    assert fleet.status_code == 200 and fleet.content.startswith(b"%PDF")

    admin = await auth_headers(client, "admin")
    archived = await client.post(f"/api/reports/archive?date={day}", headers=admin)
    assert archived.status_code == 200 and archived.json()["written"] == 2
    again = await client.post(f"/api/reports/archive?date={day}", headers=admin)
    assert again.json()["written"] == 2

    archive = (await client.get("/api/reports/archive", headers=headers)).json()
    assert len([item for item in archive if item["report_date"] == day]) == 2
    stored = await client.get(f"/api/reports/archive/{archive[0]['id']}/pdf", headers=headers)
    assert stored.status_code == 200 and stored.content.startswith(b"%PDF")


async def test_report_for_unknown_vehicle_is_404(client):
    headers = await auth_headers(client, "admin")
    response = await client.get(f"/api/reports/route-day?vehicle_id=9999&date={_yesterday()}", headers=headers)
    assert response.status_code == 404


async def test_company_locations_and_default_depot(client):
    headers = await auth_headers(client, "admin")
    created = await client.post(
        "/api/company/locations",
        json={"name": "Çivril Şube", "kind": "Şube", "latitude": 38.30, "longitude": 29.74, "is_default_depot": True},
        headers=headers,
    )
    assert created.status_code == 201, created.text
    locations = (await client.get("/api/company/locations", headers=headers)).json()
    defaults = [item for item in locations if item["is_default_depot"]]
    assert [item["name"] for item in defaults] == ["Çivril Şube"]

    config = (await client.get("/api/map/config", headers=headers)).json()
    assert (config["depot_lat"], config["depot_lng"], config["depot_label"]) == (38.30, 29.74, "Çivril Şube")

    data = (await client.get("/api/map/data", headers=headers)).json()
    assert any(item["name"] == "Çivril Şube" for item in data["company_locations"])

    original = next(item for item in locations if item["name"] == "Denizli Merkez Depo")
    restored = await client.put(
        f"/api/company/locations/{original['id']}", json={"is_default_depot": True}, headers=headers
    )
    assert restored.status_code == 200
    await client.delete(f"/api/company/locations/{created.json()['id']}", headers=headers)


async def test_entity_can_be_moved_and_new_types_listed(client):
    headers = await auth_headers(client, "admin")
    competitors = (await client.get("/api/entities?entity_type=competitor", headers=headers)).json()
    assert [item["name"] for item in competitors] == ["Ege Dağıtım (Rakip)"]
    target = competitors[0]
    moved = await client.put(
        f"/api/entities/{target['id']}", json={"latitude": 37.7711, "longitude": 29.0577}, headers=headers
    )
    assert moved.status_code == 200
    assert (moved.json()["latitude"], moved.json()["longitude"]) == (37.7711, 29.0577)
