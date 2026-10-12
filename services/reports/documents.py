"""PDF and Excel output for daily route reports."""

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from backend.config import ROOT_DIR, settings
from services.reports.route_day import RouteDay

FONT_DIR = ROOT_DIR / "backend" / "assets" / "fonts"
FONT = "DejaVuSans"
FONT_BOLD = "DejaVuSans-Bold"
INK = colors.HexColor("#122033")
PINE = colors.HexColor("#1f6f5b")


@dataclass
class RouteDayDocument:
    company_name: str
    logo_file: Path | None
    vehicle_plate: str
    vehicle_type: str
    day: date
    route: RouteDay
    drivers: list[str] = field(default_factory=list)
    map_png: bytes | None = None


def _register_fonts() -> None:
    if FONT in pdfmetrics.getRegisteredFontNames():
        return
    pdfmetrics.registerFont(TTFont(FONT, str(FONT_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, str(FONT_DIR / "DejaVuSans-Bold.ttf")))


def _clock(value: datetime | None) -> str:
    return value.astimezone(settings.tz).strftime("%H:%M") if value else "—"


def _duration(minutes: float) -> str:
    hours, rest = divmod(int(round(minutes)), 60)
    return f"{hours} sa {rest} dk" if hours else f"{rest} dk"


def _summary_rows(doc: RouteDayDocument) -> list[tuple[str, str]]:
    route = doc.route
    return [
        ("Toplam mesafe", f"{route.distance_km:.1f} km"),
        ("Sürüş süresi", _duration(route.driving_min)),
        ("Durak süresi", _duration(route.stopped_min)),
        ("Durak sayısı", str(len(route.stops))),
        ("İlk hareket", _clock(route.first_move)),
        ("Son hareket", _clock(route.last_move)),
        ("En yüksek hız", f"{route.max_speed_kmh:.0f} km/sa"),
        ("Konum kaydı", str(len(route.points))),
        ("Sürücü", ", ".join(doc.drivers) or "—"),
    ]


def _stop_rows(route: RouteDay) -> list[list[str]]:
    rows = []
    for stop in route.stops:
        if stop.place:
            place = f"{stop.place.name} ({stop.place.kind})"
        else:
            place = f"{stop.latitude:.5f}, {stop.longitude:.5f}"
        rows.append(
            [
                str(stop.index),
                _clock(stop.start),
                _clock(stop.end),
                f"{stop.duration_min:.0f}",
                place,
                "Evet" if stop.planned else "—",
            ]
        )
    return rows


def _styles() -> dict[str, ParagraphStyle]:
    return {
        "title": ParagraphStyle("title", fontName=FONT_BOLD, fontSize=16, textColor=INK, spaceAfter=4),
        "sub": ParagraphStyle("sub", fontName=FONT, fontSize=10, textColor=colors.HexColor("#4b5563")),
        "h2": ParagraphStyle("h2", fontName=FONT_BOLD, fontSize=12, textColor=INK, spaceBefore=10, spaceAfter=6),
        "body": ParagraphStyle("body", fontName=FONT, fontSize=9, textColor=INK, leading=12),
        "cell": ParagraphStyle("cell", fontName=FONT, fontSize=8.5, textColor=INK, leading=11),
    }


def _table(rows: list[list], widths: list[float], header: bool = True) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    style = [
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("TEXTCOLOR", (0, 0), (-1, -1), INK),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d1d5db")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if header:
        style += [
            ("BACKGROUND", (0, 0), (-1, 0), PINE),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ]
    table.setStyle(TableStyle(style))
    return table


def _vehicle_section(doc: RouteDayDocument, styles: dict[str, ParagraphStyle]) -> list:
    story: list = []
    header_cells: list = []
    if doc.logo_file and doc.logo_file.exists():
        try:
            header_cells.append(Image(str(doc.logo_file), width=2.6 * cm, height=2.6 * cm, kind="proportional"))
        except OSError:
            header_cells.append("")
    else:
        header_cells.append("")
    title = [
        Paragraph(doc.company_name, styles["sub"]),
        Paragraph("Günlük Rota Raporu", styles["title"]),
        Paragraph(
            f"{doc.vehicle_plate} · {doc.vehicle_type} · {doc.day.strftime('%d.%m.%Y')}",
            styles["sub"],
        ),
    ]
    header_cells.append(title)
    header = Table([header_cells], colWidths=[3 * cm, 14 * cm])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story += [header, Spacer(1, 8)]

    if not doc.route.points:
        story.append(Paragraph("Bu gün için araçtan konum verisi gelmedi.", styles["body"]))
        if doc.route.missed_planned:
            story.append(Paragraph("Planlı ama ziyaret edilmeyen: " + ", ".join(doc.route.missed_planned), styles["body"]))
        return story

    if doc.map_png:
        story.append(Image(io.BytesIO(doc.map_png), width=17 * cm, height=17 * cm * 560 / 1000))
        story.append(Spacer(1, 6))

    story.append(Paragraph("Özet", styles["h2"]))
    summary = [[label, value] for label, value in _summary_rows(doc)]
    story.append(_table(summary, [5 * cm, 12 * cm], header=False))

    story.append(Paragraph("Duraklar", styles["h2"]))
    stop_rows = _stop_rows(doc.route)
    if stop_rows:
        rows = [["#", "Varış", "Ayrılış", "Süre (dk)", "Yer", "Planlı"]]
        rows += [[*row[:4], Paragraph(row[4], styles["cell"]), row[5]] for row in stop_rows]
        story.append(_table(rows, [0.8 * cm, 1.6 * cm, 1.6 * cm, 1.8 * cm, 9.6 * cm, 1.6 * cm]))
    else:
        story.append(Paragraph("5 dakikadan uzun durak tespit edilmedi.", styles["body"]))

    if doc.route.missed_planned:
        story.append(Paragraph("Planlı ama ziyaret edilmeyen", styles["h2"]))
        story.append(Paragraph(", ".join(doc.route.missed_planned), styles["body"]))
    return story


def _footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont(FONT, 7.5)
    canvas.setFillColor(colors.HexColor("#6b7280"))
    canvas.drawString(2 * cm, 1.2 * cm, f"{settings.app_name} · Konumlar filo takip sisteminden alınmıştır")
    canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Sayfa {document.page}")
    canvas.restoreState()


def _build_pdf(story: list, title: str) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.8 * cm,
        title=title,
        author=settings.app_name,
    )
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


def route_day_pdf(doc: RouteDayDocument) -> bytes:
    _register_fonts()
    styles = _styles()
    return _build_pdf(_vehicle_section(doc, styles), f"Rota raporu {doc.vehicle_plate} {doc.day.isoformat()}")


def fleet_day_pdf(company_name: str, day: date, docs: list[RouteDayDocument]) -> bytes:
    _register_fonts()
    styles = _styles()
    story: list = [
        Paragraph(company_name, styles["sub"]),
        Paragraph("Filo Günlük Rota Özeti", styles["title"]),
        Paragraph(day.strftime("%d.%m.%Y"), styles["sub"]),
        Spacer(1, 10),
    ]
    rows = [["Plaka", "Mesafe", "Sürüş", "Durak", "İlk hareket", "Son hareket"]]
    for doc in docs:
        route = doc.route
        rows.append(
            [
                doc.vehicle_plate,
                f"{route.distance_km:.1f} km",
                _duration(route.driving_min),
                str(len(route.stops)),
                _clock(route.first_move),
                _clock(route.last_move),
            ]
        )
    if len(rows) == 1:
        story.append(Paragraph("Bu gün için takip edilen araç verisi yok.", styles["body"]))
    else:
        story.append(_table(rows, [3.2 * cm, 2.6 * cm, 3 * cm, 2 * cm, 3 * cm, 3 * cm]))
    for doc in docs:
        story.append(PageBreak())
        story += _vehicle_section(doc, styles)
    return _build_pdf(story, f"Filo raporu {day.isoformat()}")


def route_day_xlsx(doc: RouteDayDocument) -> bytes:
    workbook = Workbook()
    bold = Font(bold=True)

    summary = workbook.active
    summary.title = "Özet"
    summary.append(["Firma", doc.company_name])
    summary.append(["Araç", doc.vehicle_plate])
    summary.append(["Tür", doc.vehicle_type])
    summary.append(["Tarih", doc.day.strftime("%d.%m.%Y")])
    for label, value in _summary_rows(doc):
        summary.append([label, value])
    if doc.route.missed_planned:
        summary.append(["Planlı ama ziyaret edilmeyen", ", ".join(doc.route.missed_planned)])
    for row in summary.iter_rows(min_col=1, max_col=1):
        row[0].font = bold
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 48

    stops = workbook.create_sheet("Duraklar")
    stops.append(["#", "Varış", "Ayrılış", "Süre (dk)", "Yer", "Tür", "Mesafe (m)", "Planlı", "Enlem", "Boylam"])
    for stop in doc.route.stops:
        stops.append(
            [
                stop.index,
                _clock(stop.start),
                _clock(stop.end),
                stop.duration_min,
                stop.place.name if stop.place else "",
                stop.place.kind if stop.place else "",
                stop.distance_m,
                "Evet" if stop.planned else "Hayır",
                round(stop.latitude, 6),
                round(stop.longitude, 6),
            ]
        )
    for cell in stops[1]:
        cell.font = bold
    stops.column_dimensions["E"].width = 32

    raw = workbook.create_sheet("Konumlar")
    raw.append(["Saat", "Enlem", "Boylam", "Hız (km/sa)", "Kontak"])
    for point in doc.route.points:
        raw.append(
            [
                point.recorded_at.astimezone(settings.tz).strftime("%H:%M:%S"),
                point.latitude,
                point.longitude,
                point.speed_kmh,
                "" if point.ignition_on is None else ("Açık" if point.ignition_on else "Kapalı"),
            ]
        )
    for cell in raw[1]:
        cell.font = bold

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
