"""Customer and supplier CRUD, bulk import and geocoding."""

import csv
import io
from datetime import datetime
from typing import Any, Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import Response
from sqlalchemy import or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from backend.config import settings
from backend.database import get_db
from backend.deps import get_current_user, require_roles
from backend.models import Entity, EntityCategory, EntityStatus, EntityType, Region, User
from backend.schemas import EntityIn, EntityOut, EntityUpdate, GeocodeOut, ImportResult
from services.geocoding import geocode_address

router = APIRouter()

MAX_IMPORT_BYTES = 5 * 1024 * 1024
HEADER_ALIASES = {
    "name": "name",
    "unvan": "name",
    "ad": "name",
    "type": "type",
    "tip": "type",
    "category": "category",
    "kategori": "category",
    "address": "address",
    "adres": "address",
    "latitude": "latitude",
    "enlem": "latitude",
    "lat": "latitude",
    "longitude": "longitude",
    "boylam": "longitude",
    "lng": "longitude",
    "phone": "phone",
    "telefon": "phone",
    "email": "email",
    "eposta": "email",
    "e-posta": "email",
    "tax_number": "tax_number",
    "vergi_no": "tax_number",
    "vergi": "tax_number",
    "current_balance": "current_balance",
    "bakiye": "current_balance",
    "status": "status",
    "durum": "status",
    "region_id": "region_id",
    "bolge_id": "region_id",
    "region_code": "region_code",
    "bolge": "region_code",
    "bölge": "region_code",
    "region": "region_code",
}
TYPE_LABELS = {
    EntityType.customer: "Müşteri",
    EntityType.supplier: "Tedarikçi",
    EntityType.competitor: "Rakip",
    EntityType.other: "Diğer",
}
EXPORT_HEADERS = [
    "Unvan", "Tip", "Kategori", "Adres", "Enlem", "Boylam",
    "Telefon", "E-posta", "Vergi No", "Bakiye", "Durum", "Bölge",
]
TYPE_ALIASES = {
    "customer": EntityType.customer,
    "müşteri": EntityType.customer,
    "musteri": EntityType.customer,
    "supplier": EntityType.supplier,
    "tedarikçi": EntityType.supplier,
    "tedarikci": EntityType.supplier,
    "competitor": EntityType.competitor,
    "rakip": EntityType.competitor,
    "other": EntityType.other,
    "diğer": EntityType.other,
    "diger": EntityType.other,
}


def _entity_out(entity: Entity) -> EntityOut:
    payload = EntityOut.model_validate(entity)
    payload.region_name = entity.region.name if entity.region else None
    return payload


async def _get_entity(db: AsyncSession, entity_id: int) -> Entity:
    entity = await db.scalar(
        select(Entity).options(selectinload(Entity.region)).where(Entity.id == entity_id)
    )
    if entity is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Cari bulunamadı")
    return entity


async def _ensure_region(db: AsyncSession, region_id: int | None) -> None:
    if region_id is None:
        return
    region = await db.get(Region, region_id)
    if region is None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Bölge bulunamadı")


def _match_enum(enum_cls: type, raw: str):
    text = raw.strip()
    if not text:
        return None
    folded = text.casefold()
    for item in enum_cls:
        if item.value.casefold() == folded or item.name.casefold() == folded:
            return item
    return None


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _optional_float(value: Any) -> float | None:
    text = _cell(value).replace(" ", "")
    if not text:
        return None
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    return float(text)


def _normalize_headers(headers: list[str]) -> list[str]:
    normalized = []
    for header in headers:
        key = _cell(header).casefold().replace(" ", "_")
        normalized.append(HEADER_ALIASES.get(key, key))
    return normalized


def _rows_from_csv(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig")
    sample = text[:2048]
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_headers = next(reader, None)
    if not raw_headers:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dosyada başlık satırı yok")
    headers = _normalize_headers(raw_headers)
    rows = []
    for raw in reader:
        if not any(_cell(cell) for cell in raw):
            continue
        rows.append({headers[index]: _cell(value) for index, value in enumerate(raw) if index < len(headers)})
    return rows


def _rows_from_xlsx(content: bytes) -> list[dict[str, str]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Excel desteği yüklü değil") from exc
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Excel dosyası okunamadı") from exc
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    raw_headers = next(iterator, None)
    if not raw_headers:
        workbook.close()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dosyada başlık satırı yok")
    headers = _normalize_headers([_cell(header) for header in raw_headers])
    rows = []
    for raw in iterator:
        if not any(_cell(cell) for cell in raw):
            continue
        rows.append({headers[index]: _cell(value) for index, value in enumerate(raw) if index < len(headers)})
    workbook.close()
    return rows


@router.get("", response_model=list[EntityOut])
async def list_entities(
    search: str | None = None,
    entity_type: EntityType | None = None,
    category: EntityCategory | None = None,
    status_filter: EntityStatus | None = Query(default=None, alias="status"),
    region_id: int | None = None,
    limit: int = Query(default=300, ge=1, le=1000),
    _: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> list[EntityOut]:
    stmt = select(Entity).options(selectinload(Entity.region))
    if search and search.strip():
        like = f"%{search.strip()}%"
        stmt = stmt.where(
            or_(
                Entity.name.ilike(like),
                Entity.address.ilike(like),
                Entity.phone.ilike(like),
                Entity.tax_number.ilike(like),
            )
        )
    if entity_type is not None:
        stmt = stmt.where(Entity.type == entity_type)
    if category is not None:
        stmt = stmt.where(Entity.category == category)
    if status_filter is not None:
        stmt = stmt.where(Entity.status == status_filter)
    if region_id is not None:
        stmt = stmt.where(Entity.region_id == region_id)
    stmt = stmt.order_by(Entity.name).limit(limit)
    entities = (await db.scalars(stmt)).all()
    return [_entity_out(entity) for entity in entities]


@router.post("", response_model=EntityOut, status_code=status.HTTP_201_CREATED)
async def create_entity(
    payload: EntityIn,
    _: User = Depends(require_roles()),
    db: AsyncSession = Depends(get_db),
) -> EntityOut:
    if (payload.latitude is None) ^ (payload.longitude is None):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Enlem ve boylam birlikte gönderilmeli")
    await _ensure_region(db, payload.region_id)
    entity = Entity(**payload.model_dump())
    if entity.latitude is None or entity.longitude is None:
        result = await geocode_address(entity.address)
        if result.latitude is not None and result.longitude is not None:
            entity.latitude = result.latitude
            entity.longitude = result.longitude
    db.add(entity)
    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Cari kaydı çakışıyor") from None
    return _entity_out(await _get_entity(db, entity.id))


@router.post("/import", response_model=ImportResult)
async def import_entities(
    file: UploadFile = File(...),
    _: User = Depends(require_roles()),
    db: AsyncSession = Depends(get_db),
) -> ImportResult:
    filename = (file.filename or "").lower()
    content = await file.read()
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Dosya 5 MB sınırını aşıyor")
    try:
        if filename.endswith(".csv"):
            rows = _rows_from_csv(content)
        elif filename.endswith(".xlsx"):
            rows = _rows_from_xlsx(content)
        else:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Yalnızca CSV veya XLSX yükleyin")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dosya UTF-8 olmalı") from exc

    regions = (await db.scalars(select(Region))).all()
    region_ids = {region.id for region in regions}
    region_codes = {region.name.casefold(): region.id for region in regions}
    region_codes.update({region.code.casefold(): region.id for region in regions})
    created = 0
    warnings: list[str] = []

    for index, row in enumerate(rows, start=2):
        name = row.get("name", "").strip()
        if not name:
            warnings.append(f"Satır {index}: unvan boş, atlandı")
            continue
        entity_type = TYPE_ALIASES.get(row.get("type", "").strip().casefold())
        category = _match_enum(EntityCategory, row.get("category", ""))
        if entity_type is None or category is None:
            warnings.append(f"Satır {index}: tip veya kategori tanınmadı")
            continue
        entity_status = _match_enum(EntityStatus, row.get("status", "")) or EntityStatus.active
        try:
            latitude = _optional_float(row.get("latitude"))
            longitude = _optional_float(row.get("longitude"))
            balance = _optional_float(row.get("current_balance")) or 0
        except ValueError:
            warnings.append(f"Satır {index}: sayısal alan okunamadı")
            continue
        if (latitude is None) ^ (longitude is None):
            warnings.append(f"Satır {index}: enlem ve boylam eksik, koordinat boş bırakıldı")
            latitude = None
            longitude = None

        region_id = None
        raw_region = row.get("region_id", "").strip()
        if raw_region:
            try:
                parsed = int(float(raw_region))
            except ValueError:
                warnings.append(f"Satır {index}: bölge numarası geçersiz")
                parsed = None
            if parsed in region_ids:
                region_id = parsed
        code = row.get("region_code", "").strip().casefold()
        if region_id is None and code:
            region_id = region_codes.get(code)
            if region_id is None:
                warnings.append(f"Satır {index}: bölge kodu bulunamadı")

        if latitude is None or longitude is None:
            result = await geocode_address(row.get("address", ""))
            if result.latitude is not None and result.longitude is not None:
                latitude = result.latitude
                longitude = result.longitude
            elif result.warning:
                warnings.append(f"Satır {index}: {result.warning}")

        db.add(
            Entity(
                name=name,
                type=entity_type,
                category=category,
                address=row.get("address", ""),
                latitude=latitude,
                longitude=longitude,
                phone=row.get("phone", ""),
                email=row.get("email", ""),
                tax_number=row.get("tax_number", ""),
                current_balance=balance,
                status=entity_status,
                region_id=region_id,
            )
        )
        created += 1

    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="İçe aktarma sırasında kayıt çakışması") from None
    return ImportResult(created=created, warnings=warnings[:50])


@router.get("/export")
async def export_entities(
    fmt: Literal["xlsx", "csv"] = Query(default="xlsx", alias="format"),
    _: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> Response:
    entities = (
        await db.scalars(select(Entity).options(selectinload(Entity.region)).order_by(Entity.name))
    ).all()
    rows = [
        [
            entity.name,
            TYPE_LABELS.get(entity.type, entity.type.value),
            entity.category.value,
            entity.address,
            entity.latitude,
            entity.longitude,
            entity.phone,
            entity.email,
            entity.tax_number,
            entity.current_balance,
            entity.status.value,
            entity.region.name if entity.region else "",
        ]
        for entity in entities
    ]
    stamp = datetime.now(settings.tz).strftime("%Y-%m-%d")
    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow(EXPORT_HEADERS)
        writer.writerows(["" if value is None else value for value in row] for row in rows)
        content = buffer.getvalue().encode("utf-8-sig")
        media_type = "text/csv; charset=utf-8"
    else:
        import openpyxl
        from openpyxl.styles import Font

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Cariler"
        sheet.append(EXPORT_HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for column, width in zip("ABCDEFGHIJKL", (32, 12, 20, 44, 11, 11, 16, 28, 14, 12, 8, 16)):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        buffer = io.BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    filename = f"cariler_{stamp}.{fmt}"
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/{entity_id}", response_model=EntityOut)
async def get_entity(
    entity_id: int,
    _: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
) -> EntityOut:
    return _entity_out(await _get_entity(db, entity_id))


@router.put("/{entity_id}", response_model=EntityOut)
async def update_entity(
    entity_id: int,
    payload: EntityUpdate,
    _: User = Depends(require_roles()),
    db: AsyncSession = Depends(get_db),
) -> EntityOut:
    entity = await _get_entity(db, entity_id)
    changes = payload.model_dump(exclude_unset=True)
    if "region_id" in changes:
        await _ensure_region(db, changes["region_id"])
    lat = changes["latitude"] if "latitude" in changes else entity.latitude
    lng = changes["longitude"] if "longitude" in changes else entity.longitude
    if (lat is None) ^ (lng is None):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Enlem ve boylam birlikte gönderilmeli")
    for key, value in changes.items():
        setattr(entity, key, value)
    if "address" in changes and "latitude" not in changes and "longitude" not in changes:
        result = await geocode_address(entity.address)
        if result.latitude is not None and result.longitude is not None:
            entity.latitude = result.latitude
            entity.longitude = result.longitude
    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Cari kaydı çakışıyor") from None
    return _entity_out(await _get_entity(db, entity.id))


@router.delete("/{entity_id}")
async def delete_entity(
    entity_id: int,
    _: User = Depends(require_roles()),
    db: AsyncSession = Depends(get_db),
) -> dict:
    entity = await _get_entity(db, entity_id)
    await db.delete(entity)
    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Bu cariye bağlı sevkiyat veya ziyaret var",
        ) from None
    return {"ok": True}


@router.post("/{entity_id}/geocode", response_model=GeocodeOut)
async def geocode_entity(
    entity_id: int,
    _: User = Depends(require_roles()),
    db: AsyncSession = Depends(get_db),
) -> GeocodeOut:
    entity = await _get_entity(db, entity_id)
    result = await geocode_address(entity.address)
    if result.latitude is not None and result.longitude is not None:
        entity.latitude = result.latitude
        entity.longitude = result.longitude
        await db.commit()
        await db.refresh(entity)
    return GeocodeOut(latitude=entity.latitude, longitude=entity.longitude, warning=result.warning)
